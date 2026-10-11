"""Tests for apply-edits (corp department 7, "Apply + verify"). Temp directories only; no real file is touched.

Run with: python3 -B test_apply_edits.py        (any directory; .docx tests skip when python-docx is missing)
          /usr/bin/python3 -B test_apply_edits.py   (the text path on the system 3.9 interpreter)

The script is exercised as the CLI it is (subprocess: exit code, stdout JSON, final stderr line) and, where a failure
has to be injected (a bad write, a backup that cannot be written, a file changed underneath), in-process through its
own functions. Rebuilt .docx files are checked through python-docx's public API or the raw XML, never through the
script's own extraction, so the check is not circular.
"""
import hashlib
import importlib.machinery
import importlib.util
import io
import json
import os
import re
import shutil
import stat
import subprocess
import sys
import tempfile
import unittest
import zipfile
from contextlib import redirect_stderr, redirect_stdout
from unittest import mock

sys.dont_write_bytecode = True
HERE = os.path.dirname(os.path.realpath(__file__))
SCRIPT = os.environ.get("APPLY_EDITS_SCRIPT") or os.path.join(HERE, "apply-edits")   # env override: mutation tests only

try:
    import docx  # noqa: F401
    from docx.enum.text import WD_BREAK
    from docx.opc.constants import CONTENT_TYPE as CT
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import XmlPart
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
    HAVE_DOCX = True
except Exception:   # the system 3.9 interpreter has no python-docx: those tests skip
    HAVE_DOCX = False

BAK_RE = re.compile(r"\.bak-\d{8}-\d{6}(-\d+)?$")
# D4: a hermetic stand-in for share-check.sh (only its BUILD_STAMPS line matters); the real file is read in one test.
STAMP = "Acme Org appeal build"
FIXTURE_DIR = tempfile.mkdtemp(prefix="apply_edits_sharecheck_")
SHARE_CHECK = os.path.join(FIXTURE_DIR, "share-check.sh")
with open(SHARE_CHECK, "w", encoding="utf-8") as _fh:
    _fh.write("#!/bin/bash\npython3 - <<'PY'\nBUILD_STAMPS = {'Acme Org reconsideration build', 'Acme Org appeal build'}   # exact\nPY\n")
os.environ["CORP_SHARE_CHECK"] = SHARE_CHECK   # every CLI and in-process run in this file reads the fixture
__import__("atexit").register(shutil.rmtree, FIXTURE_DIR, True)


def load_module():
    loader = importlib.machinery.SourceFileLoader("apply_edits_under_test", SCRIPT)
    spec = importlib.util.spec_from_loader("apply_edits_under_test", loader)
    mod = importlib.util.module_from_spec(spec)
    loader.exec_module(mod)
    return mod


MOD = load_module()


def sha(path):
    with open(path, "rb") as fh:
        return hashlib.sha256(fh.read()).hexdigest()


def read(path):
    with open(path, "rb") as fh:
        return fh.read()


class Base(unittest.TestCase):
    maxDiff = None

    def setUp(self):
        self.tmp = os.path.realpath(tempfile.mkdtemp(prefix="apply_edits_test_"))
        self.src = os.path.join(self.tmp, "src")
        os.makedirs(self.src)
        self.out = os.path.join(self.tmp, "v2")

    def tearDown(self):
        for dp, dns, fns in os.walk(self.tmp):
            for n in dns + fns:
                p = os.path.join(dp, n)
                if not os.path.islink(p):
                    try:
                        os.chmod(p, 0o700 if os.path.isdir(p) else 0o600)
                    except OSError:
                        pass
        shutil.rmtree(self.tmp, ignore_errors=True)

    def put(self, rel, content, root=None):
        path = os.path.join(root or self.src, rel)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "wb") as fh:
            fh.write(content.encode("utf-8") if isinstance(content, str) else content)
        return path

    def edit_list(self, rows, bare=False, name="edits.json", **top):
        path = os.path.join(self.tmp, name)
        body = rows if bare else dict({"edits": rows, "deferred": [], "rejected": []}, **top)
        with open(path, "w", encoding="utf-8") as fh:
            json.dump(body, fh, ensure_ascii=False)
        return path

    def cli(self, *args, cwd=None):
        env = dict(os.environ, PYTHONIOENCODING="utf-8")
        r = subprocess.run([sys.executable, "-B", SCRIPT] + list(args), capture_output=True, cwd=cwd or self.tmp,
                           env=env, timeout=120)
        out, err = r.stdout.decode("utf-8"), r.stderr.decode("utf-8")
        return r.returncode, (json.loads(out) if out.strip() else {}), err

    def inproc(self, *args):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = MOD.main(list(args))
        return code, json.loads(out.getvalue()), err.getvalue()

    def last(self, err):
        return err.strip().splitlines()[-1]

    def tree(self, root):
        snap = {}
        for dp, dns, fns in os.walk(root):
            for n in fns:
                p = os.path.join(dp, n)
                snap[os.path.relpath(p, root)] = sha(p)
        return snap

    def miss(self, report, eid):
        rows = [m for m in report["misses"] if m["id"] == eid]
        self.assertEqual(len(rows), 1, report["misses"])
        return rows[0]


# ---------------------------------------------------------------------------------------------------------------
class TextApply(Base):
    def test_applied_out_dir(self):
        src = self.put("letter.md", "The quick brown fox.\nAuthorized: 30 hours/week.\n")
        before = read(src)
        el = self.edit_list([{"id": "E1", "closes": ["F1"], "file": src, "old": "30 hours", "new": "10 hours",
                              "rationale": "payer letter p2", "kind": "mechanical", "jev_agree": "agree",
                              "jev_margin": 2.0, "jev_reread": "none"}])
        code, rep, err = self.cli(el, "--out-dir", self.out)
        self.assertEqual(code, 0, err)
        self.assertEqual(self.last(err), "MISSES (0)")
        self.assertEqual((rep["ok"], rep["applied"], rep["misses"], rep["files_changed"]), (True, ["E1"], [], [src]))
        dest = os.path.join(self.out, "letter.md")
        self.assertEqual(read(dest), b"The quick brown fox.\nAuthorized: 10 hours/week.\n")
        self.assertEqual(read(src), before)
        row = rep["outputs"][0]
        self.assertEqual((row["output"], row["status"]), (dest, "written"))
        self.assertEqual((row["md5_before"], row["md5_after"]),
                         (hashlib.md5(before).hexdigest(), hashlib.md5(read(dest)).hexdigest()))

    def test_bare_list_relative_file_and_root(self):
        self.put("sub/a.txt", "alpha beta\n")
        el = self.edit_list([{"id": "E1", "file": "sub/a.txt", "old": "beta", "new": "gamma"}], bare=True)
        code, rep, err = self.cli(el, "--root", self.src, "--out-dir", self.out)
        self.assertEqual(code, 0, err)
        self.assertEqual(read(os.path.join(self.out, "sub", "a.txt")), b"alpha gamma\n")

    def test_not_found_is_a_miss_and_nothing_is_written(self):
        src = self.put("a.md", "hello world\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "hello", "new": "hi"},
                             {"id": "E2", "file": src, "old": "planet", "new": "globe"}])
        code, rep, err = self.cli(el, "--out-dir", self.out)
        self.assertEqual(code, 1)
        self.assertEqual(self.last(err), "MISSES (1)")
        self.assertEqual(self.miss(rep, "E2")["reason"], "not found")
        self.assertEqual((rep["applied"], rep["matched"], rep["written"]), ([], ["E1"], False))
        self.assertFalse(os.path.exists(self.out), "atomic: a miss must leave no output behind")

    def test_found_twice_with_default_expect_is_a_miss(self):
        src = self.put("a.md", "cat\ndog\ncat\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "cat", "new": "cow"}])
        code, rep, err = self.cli(el, "--out-dir", self.out)
        self.assertEqual(code, 1)
        m = self.miss(rep, "E1")
        self.assertEqual(m["reason"], "found 2 occurrences != expect 1")
        self.assertIn("line 1 col 1", m["hint"])
        self.assertIn("line 3 col 1", m["hint"])

    def test_expect_two_replaces_both(self):
        src = self.put("a.md", "cat\ndog\ncat\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "cat", "new": "cow", "expect": 2}])
        code, rep, err = self.cli(el, "--out-dir", self.out)
        self.assertEqual(code, 0, err)
        self.assertEqual(read(os.path.join(self.out, "a.md")), b"cow\ndog\ncow\n")

    def test_expect_higher_than_file_supports(self):
        src = self.put("a.md", "cat cat\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "cat", "new": "cow", "expect": 3}])
        code, rep, err = self.cli(el, "--dry-run")
        self.assertEqual(code, 1)
        self.assertEqual(self.miss(rep, "E1")["reason"], "found 2 occurrences != expect 3")

    def test_overlap_earlier_edit_changed_the_text(self):
        src = self.put("a.md", "The quick brown fox jumps.\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "quick brown", "new": "slow red"},
                             {"id": "E2", "file": src, "old": "brown fox", "new": "brown dog"}])
        code, rep, err = self.cli(el, "--out-dir", self.out)
        self.assertEqual(code, 1)
        m = self.miss(rep, "E2")
        self.assertTrue(m["reason"].startswith("overlaps an earlier edit: E1 changed text"), m)
        self.assertFalse(os.path.exists(self.out))

    def test_overlap_match_falls_on_text_an_earlier_edit_wrote(self):
        src = self.put("a.md", "alpha omega\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "alpha beta"},
                             {"id": "E2", "file": src, "old": "beta omega", "new": "delta"}])
        code, rep, err = self.cli(el, "--dry-run")
        self.assertEqual(code, 1)
        self.assertIn("text written by E1", self.miss(rep, "E2")["reason"])

    def test_overlap_straddling_an_earlier_deletion(self):
        src = self.put("a.md", "A B C\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": " B", "new": ""},
                             {"id": "E2", "file": src, "old": "A C", "new": "X"}])
        code, rep, err = self.cli(el, "--dry-run")
        self.assertEqual(code, 1)
        self.assertIn("text written by E1", self.miss(rep, "E2")["reason"])

    def test_old_only_inside_an_earlier_new_names_that_edit(self):
        src = self.put("a.md", "Client\u00a0name here\n")      # E1 misses (NBSP), so E2's target never exists
        el = self.edit_list([{"id": "E1", "file": src, "old": "Client name", "new": "Patient name"},
                             {"id": "E2", "file": src, "old": "Patient", "new": "Pt"}])
        code, rep, err = self.cli(el, "--dry-run")
        self.assertEqual(code, 1)
        m = self.miss(rep, "E2")
        self.assertEqual(m["reason"], "not found")
        self.assertIn("earlier edit E1", m["hint"])
        src2 = self.put("b.md", "Client name here\n")            # E1 applies: E2 then matches E1's own text
        el2 = self.edit_list([{"id": "E1", "file": src2, "old": "Client", "new": "Patient"},
                              {"id": "E2", "file": src2, "old": "Patient", "new": "Pt"}], name="e2.json")
        code, rep, err = self.cli(el2, "--dry-run")
        self.assertEqual(code, 1)
        self.assertIn("text written by E1", self.miss(rep, "E2")["reason"])

    def test_adjacent_edits_both_apply(self):
        src = self.put("a.md", "foobar\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "foo", "new": "F"},
                             {"id": "E2", "file": src, "old": "bar", "new": "B"}])
        code, rep, err = self.cli(el, "--out-dir", self.out)
        self.assertEqual(code, 0, err)
        self.assertEqual(read(os.path.join(self.out, "a.md")), b"FB\n")

    def test_list_order_sequencing_is_applied_with_a_warning(self):
        src = self.put("a.md", "X one. X two.\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "X one", "new": "Y one"},
                             {"id": "E2", "file": src, "old": "X", "new": "Z"}])
        code, rep, err = self.cli(el, "--out-dir", self.out)
        self.assertEqual(code, 0, err)
        self.assertEqual(read(os.path.join(self.out, "a.md")), b"Y one. Z two.\n")
        self.assertTrue(any(w.startswith("E2: the original had 2") for w in rep["warnings"]), rep["warnings"])

    def test_self_overlapping_occurrences_are_ambiguous(self):
        src = self.put("a.md", "aaa\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "aa", "new": "b"}])
        code, rep, err = self.cli(el, "--dry-run")
        self.assertEqual(code, 1)
        self.assertTrue(self.miss(rep, "E1")["reason"].startswith("ambiguous"))

    def test_no_op_edit_is_a_miss(self):
        src = self.put("a.md", "same\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "same", "new": "same"}])
        code, rep, err = self.cli(el, "--dry-run")
        self.assertEqual(code, 1)
        self.assertTrue(self.miss(rep, "E1")["reason"].startswith("no-op"))

    def test_judgment_kind_is_never_applied(self):
        src = self.put("a.md", "30 hours\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "30", "new": "40", "kind": "judgment"}])
        code, rep, err = self.cli(el, "--out-dir", self.out)
        self.assertEqual(code, 1)
        self.assertIn("judgment", self.miss(rep, "E1")["reason"])
        self.assertFalse(os.path.exists(self.out))

    def test_file_missing_is_a_miss(self):
        el = self.edit_list([{"id": "E1", "file": os.path.join(self.src, "nope.md"), "old": "a", "new": "b"}])
        code, rep, err = self.cli(el, "--dry-run")
        self.assertEqual(code, 1)
        self.assertTrue(self.miss(rep, "E1")["reason"].startswith("file missing"))

    def test_unsupported_binary_and_non_utf8_files_are_misses(self):
        pdf = self.put("a.pdf", b"%PDF-1.7 text")
        nul = self.put("b.bin", b"abc\x00def")
        latin = self.put("c.txt", "caf\u00e9".encode("latin-1"))
        el = self.edit_list([{"id": "E1", "file": pdf, "old": "text", "new": "x"},
                             {"id": "E2", "file": nul, "old": "abc", "new": "x"},
                             {"id": "E3", "file": latin, "old": "caf", "new": "x"}])
        code, rep, err = self.cli(el, "--dry-run")
        self.assertEqual(code, 1)
        self.assertIn("unsupported file type .pdf", self.miss(rep, "E1")["reason"])
        self.assertIn("binary file", self.miss(rep, "E2")["reason"])
        self.assertIn("not UTF-8 text", self.miss(rep, "E3")["reason"])

    def test_deletion_edit(self):
        src = self.put("a.md", "keep [DRAFT] this\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "[DRAFT] ", "new": ""}])
        code, rep, err = self.cli(el, "--out-dir", self.out)
        self.assertEqual(code, 0, err)
        self.assertEqual(read(os.path.join(self.out, "a.md")), b"keep this\n")

    def test_unicode_is_matched_exactly_with_a_hint(self):
        src = self.put("a.md", "the client\u2019s plan\n")
        before = read(src)
        el = self.edit_list([{"id": "E1", "file": src, "old": "client's plan", "new": "client's program"}])
        code, rep, err = self.cli(el, "--out-dir", self.out)
        self.assertEqual(code, 1)
        m = self.miss(rep, "E1")
        self.assertEqual(m["reason"], "not found")
        self.assertIn("U+2019 RIGHT SINGLE QUOTATION MARK", m["hint"])
        self.assertIn("NOT applied", m["hint"])
        self.assertEqual(read(src), before)
        self.assertFalse(os.path.exists(self.out))

    def test_whitespace_is_matched_exactly_with_a_hint(self):
        src = self.put("a.md", "Client\u00a0name:  A.B.\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "Client name", "new": "Patient name"},
                             {"id": "E2", "file": src, "old": "name: A.B.", "new": "name: A.B"}])
        code, rep, err = self.cli(el, "--dry-run")
        self.assertEqual(code, 1)
        h1, h2 = self.miss(rep, "E1")["hint"], self.miss(rep, "E2")["hint"]
        self.assertIn("near miss after whitespace normalisation", h1)
        self.assertIn("U+00A0 NO-BREAK SPACE", h1)
        self.assertIn("\\u00a0", h1)
        self.assertIn("near miss after whitespace normalisation", h2)
        self.assertEqual(rep["matched"], [])

    def test_case_near_miss_hint(self):
        src = self.put("a.md", "Care Plan\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "care plan", "new": "x"}])
        code, rep, err = self.cli(el, "--dry-run")
        self.assertIn("letter-case", self.miss(rep, "E1")["hint"])

    def test_longest_prefix_hint(self):
        src = self.put("a.md", "Authorized units for 10003 are 1,700 per period.\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "Authorized units for 10003 were 1,700", "new": "x"}])
        code, rep, err = self.cli(el, "--dry-run")
        hint = self.miss(rep, "E1")["hint"]
        self.assertIn("longest prefix of old found in the file is 27 of 37", hint)
        self.assertIn("U+0077 LATIN SMALL LETTER W", hint)

    def test_crlf_is_preserved_and_lf_old_misses(self):
        src = self.put("a.txt", b"line one\r\nline two\r\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "one", "new": "1"}])
        code, rep, err = self.cli(el, "--out-dir", self.out)
        self.assertEqual(code, 0, err)
        self.assertEqual(read(os.path.join(self.out, "a.txt")), b"line 1\r\nline two\r\n")
        el2 = self.edit_list([{"id": "E2", "file": src, "old": "one\nline", "new": "x"}], name="e2.json")
        code, rep, err = self.cli(el2, "--dry-run")
        self.assertEqual(code, 1)
        self.assertIn("U+000D", self.miss(rep, "E2")["hint"])

    def test_atomic_two_files_one_miss_writes_neither(self):
        a = self.put("a.md", "alpha\n")
        b = self.put("b.md", "beta\n")
        el = self.edit_list([{"id": "E1", "file": a, "old": "alpha", "new": "ALPHA"},
                             {"id": "E2", "file": b, "old": "gamma", "new": "GAMMA"}])
        code, rep, err = self.cli(el, "--out-dir", self.out)
        self.assertEqual(code, 1)
        self.assertEqual((rep["written"], rep["applied"], rep["files_changed"]), (False, [], []))
        self.assertFalse(os.path.exists(self.out))
        self.assertIn("NOTHING written", err)

    def test_out_dir_leaves_originals_byte_identical(self):
        self.put("a.md", "alpha\n")
        self.put("deep/b.md", "beta\n")
        if HAVE_DOCX:
            d = docx.Document()
            d.add_paragraph("gamma paragraph")
            d.save(os.path.join(self.src, "c.docx"))
        rows = [{"id": "E1", "file": os.path.join(self.src, "a.md"), "old": "alpha", "new": "ALPHA"},
                {"id": "E2", "file": os.path.join(self.src, "deep/b.md"), "old": "beta", "new": "BETA"}]
        if HAVE_DOCX:
            rows.append({"id": "E3", "file": os.path.join(self.src, "c.docx"), "old": "gamma", "new": "GAMMA"})
        before = self.tree(self.src)
        code, rep, err = self.cli(self.edit_list(rows), "--out-dir", self.out, "--stamp", STAMP)
        self.assertEqual(code, 0, err)
        self.assertEqual(self.tree(self.src), before, "originals (and their folders) must be byte-identical")
        self.assertEqual(read(os.path.join(self.out, "a.md")), b"ALPHA\n")
        self.assertEqual(read(os.path.join(self.out, "deep", "b.md")), b"BETA\n")
        if HAVE_DOCX:
            self.assertEqual(docx.Document(os.path.join(self.out, "c.docx")).paragraphs[0].text, "GAMMA paragraph")
        self.assertEqual(sorted(os.listdir(self.out)), sorted(["a.md", "deep"] + (["c.docx"] if HAVE_DOCX else [])))

    def test_out_dir_existing_copy_is_backed_up_not_overwritten(self):
        src = self.put("a.md", "alpha\n")
        os.makedirs(self.out)
        stale = self.put("a.md", "previous build\n", root=self.out)
        el = self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"}])
        code, rep, err = self.cli(el, "--out-dir", self.out)
        self.assertEqual(code, 0, err)
        self.assertEqual(read(stale), b"ALPHA\n")
        self.assertEqual(len(rep["backups"]), 1)
        self.assertRegex(rep["backups"][0]["backup"], BAK_RE)
        self.assertEqual(read(rep["backups"][0]["backup"]), b"previous build\n")

    def test_out_dir_that_is_the_source_dir_is_bad_input(self):
        src = self.put("a.md", "alpha\n")
        before = read(src)
        el = self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"}])
        code, rep, err = self.cli(el, "--out-dir", self.src)
        self.assertEqual(code, 2)
        self.assertTrue(self.last(err).startswith("BAD INPUT:"), err)
        self.assertNotIn("MISSES (0)", err)
        self.assertEqual(read(src), before)

    def test_dry_run_with_out_dir_writes_nothing(self):
        src = self.put("a.md", "alpha\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"}])
        code, rep, err = self.cli(el, "--out-dir", self.out, "--dry-run")
        self.assertEqual(code, 0, err)
        self.assertEqual((rep["dry_run"], rep["written"], rep["applied"], rep["matched"]), (True, False, [], ["E1"]))
        self.assertEqual(rep["outputs"][0]["status"], "planned")
        self.assertFalse(os.path.exists(self.out))

    def test_check_mode_alone(self):
        src = self.put("a.md", "alpha\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"}])
        code, rep, err = self.cli(el, "--dry-run")
        self.assertEqual((code, rep["mode"], rep["matched"]), (0, "check", ["E1"]))
        self.assertEqual(read(src), b"alpha\n")

    def test_bad_input_exits_2_and_never_prints_misses_zero(self):
        src = self.put("a.md", "alpha\n")
        cases = {
            "not json": "{nope",
            "no edits key": json.dumps({"deferred": []}),
            "missing old": json.dumps({"edits": [{"id": "E1", "file": src, "new": "x"}]}),
            "empty old": json.dumps({"edits": [{"id": "E1", "file": src, "old": "", "new": "x"}]}),
            "non-string new": json.dumps({"edits": [{"id": "E1", "file": src, "old": "a", "new": 5}]}),
            "duplicate id": json.dumps({"edits": [{"id": "E1", "file": src, "old": "a", "new": "b"},
                                                  {"id": "E1", "file": src, "old": "l", "new": "L"}]}),
            "expect zero": json.dumps({"edits": [{"id": "E1", "file": src, "old": "a", "new": "b", "expect": 0}]}),
            "expect bool": json.dumps({"edits": [{"id": "E1", "file": src, "old": "a", "new": "b", "expect": True}]}),
            "bad kind": json.dumps({"edits": [{"id": "E1", "file": src, "old": "a", "new": "b", "kind": "maybe"}]}),
            "lone surrogate": '{"edits": [{"id": "E1", "file": "%s", "old": "a", "new": "\\ud800"}]}' % src,
        }
        for name, text in cases.items():
            path = os.path.join(self.tmp, "bad.json")
            with open(path, "w", encoding="utf-8") as fh:
                fh.write(text)
            code, rep, err = self.cli(path, "--out-dir", self.out)
            self.assertEqual(code, 2, name)
            self.assertTrue(self.last(err).startswith("BAD INPUT:"), (name, err))
            self.assertNotIn("MISSES (", err, name)
            self.assertFalse(rep["ok"], name)
        good = self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"}])
        for args in (["--out-dir", self.out, "--in-place"], [], ["--backup-dir", self.tmp]):
            code, rep, err = self.cli(good, *args)
            self.assertEqual(code, 2, args)
        code, rep, err = self.cli(good, "--root", self.tmp + "/src/..", "--dry-run")
        self.assertEqual(code, 0, err)
        outside = self.edit_list([{"id": "E1", "file": "/etc/hosts", "old": "a", "new": "b"}], name="o.json")
        code, rep, err = self.cli(outside, "--root", self.src, "--dry-run")
        self.assertEqual(code, 2)
        self.assertIn("outside --root", err)
        self.assertEqual(read(src), b"alpha\n")

    def test_empty_edit_list_is_green_with_a_warning(self):
        code, rep, err = self.cli(self.edit_list([]), "--out-dir", self.out)
        self.assertEqual(code, 0, err)
        self.assertIn("the edit list has no edits; nothing to apply", rep["warnings"])
        self.assertFalse(os.path.exists(self.out))

    def test_stdin_edit_list(self):
        src = self.put("a.md", "alpha\n")
        env = dict(os.environ, PYTHONIOENCODING="utf-8")
        body = json.dumps([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"}]).encode()
        r = subprocess.run([sys.executable, "-B", SCRIPT, "-", "--dry-run"], input=body, capture_output=True, env=env)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(json.loads(r.stdout.decode())["matched"], ["E1"])

    def test_two_spellings_of_one_file_are_one_file(self):
        src = self.put("a.md", "alpha beta\n")
        link = os.path.join(self.tmp, "link.md")
        os.symlink(src, link)
        el = self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"},
                             {"id": "E2", "file": link, "old": "beta", "new": "BETA"}])
        code, rep, err = self.cli(el, "--in-place")
        self.assertEqual(code, 0, err)
        self.assertEqual(read(src), b"ALPHA BETA\n", "both edits must land; the second must not clobber the first")
        self.assertTrue(os.path.islink(link), "the symlink itself must survive an in-place edit")
        self.assertEqual(len(rep["backups"]), 1)

    def test_post_check_bites_when_the_written_bytes_are_wrong(self):
        """Gate-bite proof: corrupt the write; the post-check must go red (exit 1), never MISSES (0)."""
        src = self.put("a.md", "alpha\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"}])
        real_write = MOD._write_atomic

        def corrupt(dest, data, mode_src):
            real_write(dest, data.replace(b"ALPHA", b"ALPHX"), mode_src)

        with mock.patch.object(MOD, "_write_atomic", corrupt):
            code, rep, err = self.inproc(el, "--out-dir", self.out)
        self.assertEqual(code, 1)
        self.assertEqual(self.last(err), "MISSES (1)")
        m = self.miss(rep, "E1")
        self.assertTrue(m["reason"].startswith("post-check:"), m)
        self.assertEqual((rep["written"], rep["applied"]), (True, []))
        # unpatched, the same run is green
        code, rep, err = self.inproc(el, "--out-dir", os.path.join(self.tmp, "v3"))
        self.assertEqual(code, 0, err)

    def test_pre_write_check_bites_when_serialisation_is_wrong(self):
        src = self.put("a.md", "alpha\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"}])
        with mock.patch.object(MOD.TextDoc, "serialize", lambda self: b"something else\n"):
            code, rep, err = self.inproc(el, "--out-dir", self.out)
        self.assertEqual(code, 1)
        self.assertIn("post-check (pre-write)", self.miss(rep, "E1")["reason"])
        self.assertFalse(os.path.exists(self.out))


# ---------------------------------------------------------------------------------------------------------------
class InPlace(Base):
    def test_in_place_creates_a_verified_dated_backup(self):
        src = self.put("a.sh", "echo alpha\n")
        os.chmod(src, 0o754)
        before = read(src)
        el = self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"}])
        code, rep, err = self.cli(el, "--in-place")
        self.assertEqual(code, 0, err)
        self.assertEqual(read(src), b"echo ALPHA\n")
        self.assertEqual(stat.S_IMODE(os.stat(src).st_mode), 0o754, "permissions must survive")
        self.assertEqual(len(rep["backups"]), 1)
        bak = rep["backups"][0]
        self.assertEqual(bak["file"], src)
        self.assertEqual(os.path.dirname(bak["backup"]), self.src)
        self.assertRegex(os.path.basename(bak["backup"]), r"^a\.sh\.bak-\d{8}-\d{6}$")
        self.assertEqual(read(bak["backup"]), before)
        self.assertEqual(rep["files_changed"], [src])

    def test_in_place_backup_dir_mirrors_paths(self):
        src = self.put("docs/a.md", "alpha\n")
        self.put("other/b.md", "beta\n")
        bdir = os.path.join(self.tmp, "backups")
        el = self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"},
                             {"id": "E2", "file": os.path.join(self.src, "other/b.md"), "old": "beta", "new": "B"}])
        code, rep, err = self.cli(el, "--in-place", "--backup-dir", bdir)
        self.assertEqual(code, 0, err)
        backups = sorted(b["backup"] for b in rep["backups"])
        self.assertEqual(len(backups), 2)
        self.assertTrue(backups[0].startswith(os.path.join(bdir, "docs", "a.md.bak-")), backups)
        self.assertTrue(backups[1].startswith(os.path.join(bdir, "other", "b.md.bak-")), backups)
        self.assertEqual(read(backups[0]), b"alpha\n")

    def test_in_place_refuses_when_backup_dir_is_a_file(self):
        src = self.put("a.md", "alpha\n")
        blocker = self.put("not-a-dir", "x", root=self.tmp)
        el = self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"}])
        code, rep, err = self.cli(el, "--in-place", "--backup-dir", blocker)
        self.assertEqual(code, 2)
        self.assertTrue(self.last(err).startswith("REFUSED:"), err)
        self.assertEqual(read(src), b"alpha\n")
        self.assertEqual(sorted(os.listdir(self.src)), ["a.md"])

    @unittest.skipIf(hasattr(os, "geteuid") and os.geteuid() == 0, "root ignores directory permissions")
    def test_in_place_refuses_when_backup_location_is_read_only(self):
        src = self.put("a.md", "alpha\n")
        locked = os.path.join(self.tmp, "locked")
        os.makedirs(locked)
        os.chmod(locked, 0o555)
        el = self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"}])
        code, rep, err = self.cli(el, "--in-place", "--backup-dir", os.path.join(locked, "bk"))
        self.assertEqual(code, 2)
        self.assertIn("never runs without one", err)
        self.assertEqual(read(src), b"alpha\n")

    def test_in_place_refuses_when_the_backup_write_fails(self):
        src = self.put("a.md", "alpha\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"}])

        def boom(data, stat_src, base_path):
            raise OSError("disk full")

        with mock.patch.object(MOD, "_write_backup", boom):
            code, rep, err = self.inproc(el, "--in-place")
        self.assertEqual(code, 2)
        self.assertTrue(self.last(err).startswith("REFUSED:"), err)
        self.assertEqual((rep["written"], rep["applied"]), (False, []))
        self.assertEqual(read(src), b"alpha\n")

    def test_backup_failing_on_the_second_file_touches_neither_original(self):
        a = self.put("a.md", "alpha\n")
        b = self.put("b.md", "beta\n")
        el = self.edit_list([{"id": "E1", "file": a, "old": "alpha", "new": "ALPHA"},
                             {"id": "E2", "file": b, "old": "beta", "new": "BETA"}])
        real_backup = MOD._write_backup
        calls = []

        def second_fails(data, stat_src, base_path):
            calls.append(base_path)
            if len(calls) == 2:
                raise OSError("disk full")
            return real_backup(data, stat_src, base_path)

        with mock.patch.object(MOD, "_write_backup", second_fails):
            code, rep, err = self.inproc(el, "--in-place")
        self.assertEqual(code, 2)
        self.assertEqual((read(a), read(b)), (b"alpha\n", b"beta\n"))
        self.assertEqual(len(rep["backups"]), 1, "the backup already made is reported, and kept")
        self.assertEqual(read(rep["backups"][0]["backup"]), b"alpha\n")
        self.assertIn("BACKUP kept", err)

    def test_in_place_never_overwrites_an_existing_backup(self):
        src = self.put("a.md", "alpha\n")
        sentinel = self.put("a.md.bak-20260928-120000", "an older backup\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"}])
        with mock.patch.object(MOD, "_timestamp", lambda: "20260928-120000"):
            code, rep, err = self.inproc(el, "--in-place")
        self.assertEqual(code, 0, err)
        self.assertEqual(read(sentinel), b"an older backup\n")
        self.assertEqual(rep["backups"][0]["backup"], sentinel + "-2")
        self.assertEqual(read(sentinel + "-2"), b"alpha\n")

    def test_in_place_refuses_when_the_file_changed_after_it_was_read(self):
        src = self.put("a.md", "alpha\n")
        el = self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"}])
        real_open = MOD.open_doc

        def open_then_someone_writes(real):
            result = real_open(real)
            with open(real, "ab") as fh:
                fh.write(b"another writer\n")
            return result

        with mock.patch.object(MOD, "open_doc", open_then_someone_writes):
            code, rep, err = self.inproc(el, "--in-place")
        self.assertEqual(code, 2)
        self.assertIn("changed on disk after it was read", err)
        self.assertEqual(read(src), b"alpha\nanother writer\n")
        self.assertEqual(sorted(os.listdir(self.src)), ["a.md"], "no backup, no temp file")

    def test_in_place_miss_changes_nothing_and_makes_no_backup(self):
        a = self.put("a.md", "alpha\n")
        b = self.put("b.md", "beta\n")
        el = self.edit_list([{"id": "E1", "file": a, "old": "alpha", "new": "ALPHA"},
                             {"id": "E2", "file": b, "old": "zeta", "new": "Z"}])
        before = self.tree(self.src)
        code, rep, err = self.cli(el, "--in-place")
        self.assertEqual(code, 1)
        self.assertEqual(self.tree(self.src), before)
        self.assertEqual(rep["backups"], [])

    @unittest.skipIf(hasattr(os, "geteuid") and os.geteuid() == 0, "root ignores file permissions")
    def test_in_place_refuses_a_read_only_file(self):
        src = self.put("a.md", "alpha\n")
        os.chmod(src, 0o444)
        el = self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"}])
        code, rep, err = self.cli(el, "--in-place")
        self.assertEqual(code, 2)
        self.assertIn("read-only", err)
        self.assertEqual(read(src), b"alpha\n")


# ---------------------------------------------------------------------------------------------------------------
MCNS = ('xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:v="urn:schemas-microsoft-com:vml"')


def run_props(paragraph):
    return [(r.text, bool(r.bold)) for r in paragraph.runs]


def zip_xml(path, name):
    with zipfile.ZipFile(path) as z:
        return z.read(name).decode("utf-8")


@unittest.skipUnless(HAVE_DOCX, "python-docx not importable under this interpreter")
class Docx(Base):
    def build(self, name="doc.docx"):
        d = docx.Document()
        p = d.add_paragraph()
        p.add_run("Hello ")
        p.add_run("bold world").bold = True
        p.add_run(" end.")
        d.add_paragraph("Name:\tJohn Smith")
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "Units authorized"
        t.cell(0, 1).text = "1,700"
        d.sections[0].header.paragraphs[0].text = "Acme header"
        d.sections[0].footer.paragraphs[0].text = "Footer v1"
        path = os.path.join(self.src, name)
        d.save(path)
        return path

    def apply(self, rows, *args):
        code, rep, err = self.cli(self.edit_list(rows), *(args or ("--out-dir", self.out, "--stamp", STAMP)))
        return code, rep, err

    def test_edit_inside_one_run_keeps_its_formatting(self):
        src = self.build()
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "bold world", "new": "strong world"}])
        self.assertEqual(code, 0, err)
        out = docx.Document(os.path.join(self.out, "doc.docx"))
        self.assertEqual(run_props(out.paragraphs[0]), [("Hello ", False), ("strong world", True), (" end.", False)])
        self.assertEqual(rep["formatting_risk"], [])

    def test_changed_text_across_runs_is_applied_with_formatting_risk(self):
        src = self.build()
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "Hello bold", "new": "Greetings, brave"}])
        self.assertEqual(code, 0, err)
        out = docx.Document(os.path.join(self.out, "doc.docx"))
        self.assertEqual(out.paragraphs[0].text, "Greetings, brave world end.")
        self.assertEqual(run_props(out.paragraphs[0]), [("Greetings, brave", False), (" world", True), (" end.", False)])
        self.assertEqual(len(rep["formatting_risk"]), 1)
        risk = rep["formatting_risk"][0]
        self.assertEqual((risk["id"], risk["runs"], risk["same_formatting"]), ("D1", 2, False))
        self.assertIn("document.xml paragraph 1", risk["location"])
        self.assertIn("RISK D1", err)

    def test_only_the_changed_text_is_rewritten(self):
        src = self.build()
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "lo bold", "new": "lo, bold"},
                                     {"id": "D2", "file": src, "old": "Smith", "new": "Dr. Smith"}])
        self.assertEqual(code, 0, err)
        out = docx.Document(os.path.join(self.out, "doc.docx"))
        self.assertEqual(run_props(out.paragraphs[0]), [("Hello, ", False), ("bold world", True), (" end.", False)])
        self.assertEqual(out.paragraphs[1].text, "Name:\tJohn Dr. Smith")
        self.assertEqual(rep["formatting_risk"], [])

    def test_table_header_footer_and_footnote(self):
        d = docx.Document()
        d.add_paragraph("Body text")
        t = d.add_table(rows=1, cols=1)
        t.cell(0, 0).text = "Cell 1,700"
        d.sections[0].header.paragraphs[0].text = "Header v1"
        d.sections[0].footer.paragraphs[0].text = "Footer v1"
        fn = parse_xml('<w:footnotes %s><w:footnote w:id="1"><w:p><w:r><w:t>Footnote source v1</w:t></w:r></w:p>'
                       '</w:footnote></w:footnotes>' % nsdecls("w"))
        part = XmlPart(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, fn, d.part.package)
        d.part.relate_to(part, RT.FOOTNOTES)
        src = os.path.join(self.src, "doc.docx")
        d.save(src)
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "Cell 1,700", "new": "Cell 2,000"},
                                     {"id": "D2", "file": src, "old": "Header v1", "new": "Header v2"},
                                     {"id": "D3", "file": src, "old": "Footer v1", "new": "Footer v2"},
                                     {"id": "D4", "file": src, "old": "source v1", "new": "source v2"}])
        self.assertEqual(code, 0, err)
        dest = os.path.join(self.out, "doc.docx")
        out = docx.Document(dest)
        self.assertEqual(out.tables[0].cell(0, 0).text, "Cell 2,000")
        self.assertEqual(out.sections[0].header.paragraphs[0].text, "Header v2")
        self.assertEqual(out.sections[0].footer.paragraphs[0].text, "Footer v2")
        self.assertRegex(zip_xml(dest, "word/footnotes.xml"), r"<w:t[^>]*>Footnote source v2</w:t>")

    def test_tab_in_old_and_line_break_in_new(self):
        src = self.build()
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "Name:\tJohn", "new": "Name:\tJane"},
                                     {"id": "D2", "file": src, "old": "Footer v1", "new": "Footer\nv2"}])
        self.assertEqual(code, 0, err)
        dest = os.path.join(self.out, "doc.docx")
        out = docx.Document(dest)
        self.assertEqual(out.paragraphs[1].text, "Name:\tJane Smith")
        self.assertEqual(out.sections[0].footer.paragraphs[0].text, "Footer\nv2")
        footer_xml = "".join(zip_xml(dest, n) for n in zipfile.ZipFile(dest).namelist() if "footer" in n)
        self.assertIn("<w:br/>", footer_xml)

    def test_old_spanning_paragraphs_misses_with_a_hint(self):
        src = self.build()
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "end.\nName:", "new": "x"}])
        self.assertEqual(code, 1)
        m = self.miss(rep, "D1")
        self.assertEqual(m["reason"], "not found")
        self.assertIn("spans a paragraph break", m["hint"])

    def test_expect_counts_across_paragraphs(self):
        d = docx.Document()
        d.add_paragraph("Client: A.B.")
        d.add_paragraph("The client A.B. attends.")
        src = os.path.join(self.src, "doc.docx")
        d.save(src)
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "A.B.", "new": "C.D.", "expect": 2}])
        self.assertEqual(code, 0, err)
        out = docx.Document(os.path.join(self.out, "doc.docx"))
        self.assertEqual([p.text for p in out.paragraphs], ["Client: C.D.", "The client C.D. attends."])

    def test_hyperlink_and_tracked_changes(self):
        d = docx.Document()
        p = d.add_paragraph("See ")
        rid = d.part.relate_to("https://example.org", RT.HYPERLINK, is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), rid)
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "the payer letter"
        r.append(t)
        h.append(r)
        p._p.append(h)
        p2 = d.add_paragraph("Kept ")
        p2._p.append(parse_xml('<w:del %s w:id="1" w:author="a"><w:r><w:delText>REMOVED</w:delText></w:r></w:del>'
                               % nsdecls("w")))
        p2._p.append(parse_xml('<w:ins %s w:id="2" w:author="a"><w:r><w:t>inserted text</w:t></w:r></w:ins>'
                               % nsdecls("w")))
        src = os.path.join(self.src, "doc.docx")
        d.save(src)
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "payer letter", "new": "payer's letter"},
                                     {"id": "D2", "file": src, "old": "inserted text", "new": "inserted words"},
                                     {"id": "D3", "file": src, "old": "REMOVED", "new": "x"}])
        self.assertEqual(code, 1)
        self.assertEqual(self.miss(rep, "D3")["reason"], "not found", "tracked-deletion text is not document text")
        self.assertEqual(rep["matched"], ["D1", "D2"])
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "payer letter", "new": "payer's letter"},
                                     {"id": "D2", "file": src, "old": "inserted text", "new": "inserted words"}])
        self.assertEqual(code, 0, err)
        xml = zip_xml(os.path.join(self.out, "doc.docx"), "word/document.xml")
        self.assertRegex(xml, r"<w:hyperlink [^>]*><w:r><w:t[^>]*>the payer's letter</w:t></w:r></w:hyperlink>")
        self.assertIn("<w:delText>REMOVED</w:delText>", xml)
        self.assertRegex(xml, r"<w:ins [^>]*><w:r><w:t[^>]*>inserted words</w:t></w:r></w:ins>")

    def test_text_box_fallback_copy_is_synced(self):
        d = docx.Document()
        p = d.add_paragraph("Outer ")
        box = ('<w:r %s %s><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wps:txbx><w:txbxContent>'
               '<w:p><w:r><w:t>Box says 30 hours</w:t></w:r></w:p></w:txbxContent></wps:txbx></w:drawing></mc:Choice>'
               '<mc:Fallback><w:pict><v:textbox><w:txbxContent><w:p><w:r><w:t>Box says 30 hours</w:t></w:r></w:p>'
               '</w:txbxContent></v:textbox></w:pict></mc:Fallback></mc:AlternateContent></w:r>') % (nsdecls("w"), MCNS)
        p._p.append(parse_xml(box))
        src = os.path.join(self.src, "doc.docx")
        d.save(src)
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "30 hours", "new": "10 hours"}])
        self.assertEqual(code, 0, err)
        xml = zip_xml(os.path.join(self.out, "doc.docx"), "word/document.xml")
        self.assertEqual(xml.count("Box says 10 hours"), 2, "visible copy and legacy fallback copy both edited")
        self.assertNotIn("30 hours", xml)
        self.assertEqual(rep["formatting_risk"], [])

    def test_field_code_inside_the_changed_span_is_reported(self):
        d = docx.Document()
        p = d.add_paragraph()
        p.add_run("Page ")
        for frag in ('<w:r %s><w:fldChar w:fldCharType="begin"/></w:r>',
                     '<w:r %s><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>',
                     '<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>',
                     '<w:r %s><w:t>3</w:t></w:r>',
                     '<w:r %s><w:fldChar w:fldCharType="end"/></w:r>'):
            p._p.append(parse_xml(frag % nsdecls("w")))
        p.add_run(" of 10")
        src = os.path.join(self.src, "doc.docx")
        d.save(src)
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "Page 3 of", "new": "Page_3_of"}])
        self.assertEqual(code, 0, err)
        self.assertEqual(docx.Document(os.path.join(self.out, "doc.docx")).paragraphs[0].text, "Page_3_of 10")
        risk = rep["formatting_risk"][0]
        self.assertGreater(risk["runs"], 1)
        self.assertIn("w:fldChar", risk["detail"])
        self.assertIn("w:instrText", risk["detail"])

    def test_page_break_inside_the_changed_span_is_kept_and_reported(self):
        d = docx.Document()
        p = d.add_paragraph()
        r = p.add_run("before")
        r.add_break(WD_BREAK.PAGE)
        r.add_text("after")
        src = os.path.join(self.src, "doc.docx")
        d.save(src)
        # an insertion beside the break: only "-" is written, the break stays where it was, nothing to report
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "beforeafter", "new": "before-after"}])
        self.assertEqual(code, 0, err)
        xml = zip_xml(os.path.join(self.out, "doc.docx"), "word/document.xml")
        self.assertRegex(xml, r'>before-</w:t><w:br w:type="page"/><w:t>after<')
        self.assertEqual(rep["formatting_risk"], [])
        # a change that spans the break: applied, the break is kept (after the new text) and reported
        code, rep, err = self.cli(self.edit_list([{"id": "D2", "file": src, "old": "beforeafter",
                                                    "new": "BEFOREAFTER"}], name="e2.json"),
                                  "--out-dir", os.path.join(self.tmp, "v3"), "--stamp", STAMP)
        self.assertEqual(code, 0, err)
        dest = os.path.join(self.tmp, "v3", "doc.docx")
        self.assertEqual(docx.Document(dest).paragraphs[0].text, "BEFOREAFTER")
        self.assertIn('<w:br w:type="page"/>', zip_xml(dest, "word/document.xml"))
        self.assertIn("w:br", rep["formatting_risk"][0]["detail"])

    def test_leading_and_trailing_spaces_are_preserved(self):
        d = docx.Document()
        d.add_paragraph("A-B")
        src = os.path.join(self.src, "doc.docx")
        d.save(src)
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "-", "new": " - "}])
        self.assertEqual(code, 0, err)
        xml = zip_xml(os.path.join(self.out, "doc.docx"), "word/document.xml")
        self.assertIn('<w:t xml:space="preserve">A - B</w:t>', xml)

    def test_in_place_docx_backup_and_word_lock(self):
        src = self.build()
        before = read(src)
        lock = self.put("~$doc.docx", b"owner")
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "Footer v1", "new": "Footer v2"}], "--in-place",
                                    "--stamp", STAMP)
        self.assertEqual(code, 2)
        self.assertIn("open in Word", err)
        self.assertEqual(read(src), before)
        os.rename(lock, os.path.join(self.tmp, "lock-moved-aside"))
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "Footer v1", "new": "Footer v2"}], "--in-place",
                                    "--stamp", STAMP)
        self.assertEqual(code, 0, err)
        self.assertEqual(docx.Document(src).sections[0].footer.paragraphs[0].text, "Footer v2")
        self.assertEqual(read(rep["backups"][0]["backup"]), before)

    def test_docx_near_miss_hint_names_the_paragraph(self):
        d = docx.Document()
        d.add_paragraph("Intro")
        d.add_paragraph("Authorized\u00a0units: 1,700")
        src = os.path.join(self.src, "doc.docx")
        d.save(src)
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "Authorized units", "new": "x"}])
        self.assertEqual(code, 1)
        hint = self.miss(rep, "D1")["hint"]
        self.assertIn("document.xml paragraph 2", hint)
        self.assertIn("U+00A0", hint)

    def test_corrupt_docx_is_a_miss(self):
        src = self.put("broken.docx", b"this is not a zip")
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "zip", "new": "x"}])
        self.assertEqual(code, 1)
        self.assertTrue(self.miss(rep, "D1")["reason"].startswith("cannot open as .docx"))

    def test_control_character_in_new_is_a_miss(self):
        src = self.build()
        code, rep, err = self.apply([{"id": "D1", "file": src, "old": "Footer v1", "new": "Footer\u000bv2"}])
        self.assertEqual(code, 1)
        self.assertIn("cannot hold", self.miss(rep, "D1")["reason"])

    def test_docx_unavailable_is_a_miss_naming_the_interpreter(self):
        src = self.build()
        el = self.edit_list([{"id": "D1", "file": src, "old": "Footer v1", "new": "Footer v2"}])
        with mock.patch.object(MOD, "_DOCX", ImportError("No module named 'docx'")):
            code, rep, err = self.inproc(el, "--out-dir", self.out, "--stamp", STAMP)
        self.assertEqual(code, 1)
        self.assertIn("python-docx is not importable under", self.miss(rep, "D1")["reason"])
        self.assertFalse(os.path.exists(self.out))


def core_props(path):
    """docProps/core.xml's lastModifiedBy and modified, read from the zip itself (never through the script)."""
    xml = zip_xml(path, "docProps/core.xml")
    lmb = re.search(r"<cp:lastModifiedBy>([^<]*)</cp:lastModifiedBy>", xml)
    mod = re.search(r"<dcterms:modified[^>]*>([^<]*)</dcterms:modified>", xml)
    return (lmb.group(1) if lmb else None), (mod.group(1) if mod else None)


# ---------------------------------------------------------------------------------------------------------------
@unittest.skipUnless(HAVE_DOCX, "python-docx not importable under this interpreter")
class Stamp(Base):
    """D4 / doctrine rule 20: every written .docx carries the build stamp, never the last human's name."""

    def human_draft(self, name="draft.docx"):
        d = docx.Document()
        d.add_paragraph("Units authorized: 1,700")
        d.core_properties.last_modified_by = "Alex Reviewer"          # her Word save
        d.core_properties.modified = __import__("datetime").datetime(2026, 9, 20, 14, 0, 0)
        path = os.path.join(self.src, name)
        d.save(path)
        return path

    def test_every_written_docx_carries_the_stamp_and_the_build_time(self):
        src = self.human_draft()
        el = self.edit_list([{"id": "D1", "file": src, "old": "1,700", "new": "2,000"}])
        code, rep, err = self.cli(el, "--out-dir", self.out, "--stamp", STAMP)
        self.assertEqual(code, 0, err)
        out = os.path.join(self.out, "draft.docx")
        who, when = core_props(out)
        self.assertEqual(who, STAMP)
        self.assertTrue(when.startswith(__import__("datetime").datetime.now(__import__("datetime").timezone.utc)
                                        .strftime("%Y-%m-%d")), when)   # the build time, not the 09-20 reviewer save
        self.assertEqual(core_props(src)[0], "Alex Reviewer")           # the reviewer's draft is never touched
        self.assertEqual(rep["stamp"], STAMP)
        self.assertEqual(rep["outputs"][0]["stamped"], True)
        self.assertEqual(rep["outputs"][0]["md5_after"], hashlib.md5(read(out)).hexdigest())   # receipt = stamped bytes

    def test_in_place_stamps_too(self):
        src = self.human_draft()
        el = self.edit_list([{"id": "D1", "file": src, "old": "1,700", "new": "2,000"}])
        code, rep, err = self.cli(el, "--in-place", "--stamp", "Acme Org reconsideration build")
        self.assertEqual(code, 0, err)
        self.assertEqual(core_props(src)[0], "Acme Org reconsideration build")
        self.assertEqual(core_props(rep["backups"][0]["backup"])[0], "Alex Reviewer")   # the backup keeps the reviewer's save

    def test_a_docx_write_without_stamp_exits_2_naming_the_choices(self):
        src = self.human_draft()
        before = read(src)
        el = self.edit_list([{"id": "D1", "file": src, "old": "1,700", "new": "2,000"}])
        for args in (("--out-dir", self.out), ("--in-place",)):
            with self.subTest(args=args):
                code, rep, err = self.cli(el, *args)
                self.assertEqual(code, 2, err)
                self.assertTrue(self.last(err).startswith("BAD INPUT:"), err)
                self.assertIn("'Acme Org reconsideration build'", err)
                self.assertIn("'Acme Org appeal build'", err)
                self.assertNotIn("MISSES (0)", err)
        self.assertEqual(read(src), before)
        self.assertFalse(os.path.exists(self.out))
        code, rep, err = self.cli(el, "--dry-run")                     # a check needs no stamp; it warns
        self.assertEqual(code, 0, err)
        self.assertTrue(any("needs --stamp" in w for w in rep["warnings"]), rep["warnings"])

    def test_stamp_values(self):
        src = self.human_draft()
        el = self.edit_list([{"id": "D1", "file": src, "old": "1,700", "new": "2,000"}])
        for bad in ("   ", "Acme Org appeal Build", "acme org assessment build", "x" * 256):
            with self.subTest(stamp=bad):
                code, rep, err = self.cli(el, "--out-dir", self.out, "--stamp", bad)
                self.assertEqual(code, 2, err)
        self.assertFalse(os.path.exists(self.out))
        code, rep, err = self.cli(el, "--out-dir", self.out, "--stamp", "Acme release build")
        self.assertEqual(code, 0, err)                                   # non-Acme: allowed, with the warning
        self.assertTrue(any("share-check will read this build as a human" in w for w in rep["warnings"]), rep)
        self.assertEqual(core_props(os.path.join(self.out, "draft.docx"))[0], "Acme release build")

    def test_text_only_list_ignores_the_stamp_with_a_warning(self):
        src = self.put("a.md", "alpha\n")
        code, rep, err = self.cli(self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "ALPHA"}]),
                                  "--out-dir", self.out, "--stamp", STAMP)
        self.assertEqual(code, 0, err)
        self.assertTrue(any("applies to .docx outputs only" in w for w in rep["warnings"]), rep["warnings"])
        code, rep, err = self.cli(self.edit_list([{"id": "E1", "file": src, "old": "alpha", "new": "BETA"}],
                                                 name="e2.json"), "--out-dir", os.path.join(self.tmp, "v3"))
        self.assertEqual(code, 0, err)                                   # and needs none

    def test_stamps_are_read_from_share_check_never_copied(self):
        old = os.environ.pop("CORP_SHARE_CHECK")
        try:
            stamps, why = MOD.build_stamps()                             # the real scripts/share-check.sh
        finally:
            os.environ["CORP_SHARE_CHECK"] = old
        self.assertIsNone(why)
        for s in ("Acme Org reconsideration build", "Acme Org assessment build", "Acme Org P2P brief build",
                  "Acme Org appeal build"):                                # FB-05: the four committed identities
            self.assertIn(s, stamps)
        with open(SCRIPT, encoding="utf-8") as fh:
            body = fh.read()
        self.assertNotIn("Acme Org reconsideration build", body)       # no hard-coded copy in apply-edits
        broken = os.path.join(self.tmp, "share-check.sh")
        with open(broken, "w") as fh:
            fh.write("#!/bin/bash\necho no stamps here\n")
        os.environ["CORP_SHARE_CHECK"] = broken
        try:
            src = self.human_draft()
            el = self.edit_list([{"id": "D1", "file": src, "old": "1,700", "new": "2,000"}])
            code, rep, err = self.cli(el, "--out-dir", self.out, "--stamp", STAMP)
            self.assertEqual(code, 2, err)                               # a Acme stamp is never written blind
            self.assertIn("could not be read", err)
        finally:
            os.environ["CORP_SHARE_CHECK"] = old


if __name__ == "__main__":
    unittest.main(verbosity=1)
