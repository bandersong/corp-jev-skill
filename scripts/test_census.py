"""Unit tests for census (CEO decisions E3, E4): every part of a .docx from raw XML, text and PDF; the strip.

Pure code: no Jev, no network. Run with: python3 -m unittest test_census -v   (from this directory)
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile

HERE = os.path.dirname(os.path.realpath(__file__))
CENSUS = os.path.join(HERE, "census")
W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
NS = (W + ' xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
      ' xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
      ' xmlns:v="urn:schemas-microsoft-com:vml"')


def para(*runs):
    return "<w:p>" + "".join(f'<w:r><w:t xml:space="preserve">{r}</w:t></w:r>' for r in runs) + "</w:p>"


def textbox(inner, fallback=None):
    fb = fallback if fallback is not None else inner
    return ("<w:p><w:r><mc:AlternateContent><mc:Choice Requires=\"wps\"><w:drawing><wps:txbx><w:txbxContent>"
            f"{inner}</w:txbxContent></wps:txbx></w:drawing></mc:Choice><mc:Fallback><w:pict><v:textbox>"
            f"<w:txbxContent>{fb}</w:txbxContent></v:textbox></w:pict></mc:Fallback></mc:AlternateContent></w:r></w:p>")


def docx(path, body="", header="", footer="", footnotes="", comments="", title="Appeal"):
    """A .docx package written from raw XML: every part census must read."""
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/'
                   'package/2006/content-types"/>')
        z.writestr("word/document.xml", f'<?xml version="1.0"?><w:document {NS}><w:body>{body}</w:body></w:document>')
        if header:
            z.writestr("word/header1.xml", f'<?xml version="1.0"?><w:hdr {NS}>{header}</w:hdr>')
        if footer:
            z.writestr("word/footer1.xml", f'<?xml version="1.0"?><w:ftr {NS}>{footer}</w:ftr>')
        if footnotes:
            z.writestr("word/footnotes.xml", f'<?xml version="1.0"?><w:footnotes {NS}><w:footnote w:id="1">'
                       f'{footnotes}</w:footnote></w:footnotes>')
        if comments:
            z.writestr("word/comments.xml", f'<?xml version="1.0"?><w:comments {NS}><w:comment w:id="0">'
                       f'{comments}</w:comment></w:comments>')
        z.writestr("word/styles.xml", f'<?xml version="1.0"?><w:styles {W}/>')
        z.writestr("docProps/core.xml", '<?xml version="1.0"?><cp:coreProperties xmlns:cp="x" xmlns:dc="d">'
                   f'<dc:title>{title}</dc:title><cp:lastModifiedBy>Acme Org appeal build</cp:lastModifiedBy>'
                   '</cp:coreProperties>')
    return path


A = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
C = 'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"'
STRICT_W = 'xmlns:w="http://purl.oclc.org/ooxml/wordprocessingml/main"'


def package(path, parts):
    """An Office package from raw parts {name: xml text or bytes} (the re-gate fixtures' shape)."""
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/'
                   'package/2006/content-types"/>')
        for name, data in parts.items():
            z.writestr(name, data)
    return path


def package_bytes(parts):
    import io
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return buf.getvalue()


def make_pdf(path, pages):
    """A minimal PDF, one Helvetica text line per page ("" = a page with no text layer, like a scan)."""
    n, objs = len(pages), ["<< /Type /Catalog /Pages 2 0 R >>"]
    objs.append("<< /Type /Pages /Kids [%s] /Count %d >>" % (" ".join(f"{3 + 2 * i} 0 R" for i in range(n)), n))
    for i, text in enumerate(pages):
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET" if text else "0 0 m"
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
                    f"/Resources << /Font << /F1 {3 + 2 * n} 0 R >> >> >>")
        objs.append(f"<< /Length {len(content)} >>\nstream\n{content}\nendstream")
    objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out, offs = "%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out.encode()))
        out += f"{i} 0 obj\n{o}\nendobj\n"
    xref = len(out.encode())
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n" + "".join(f"{o:010d} 00000 n \n" for o in offs)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    with open(path, "w") as fh:
        fh.write(out)
    return path


class CensusCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="census_test_")
        self.addCleanup(shutil.rmtree, self.tmp, True)

    def p(self, name):
        return os.path.join(self.tmp, name)

    def census(self, *args):
        r = subprocess.run([sys.executable, CENSUS, *args], capture_output=True, text=True, timeout=120)
        lines = [json.loads(l) for l in r.stdout.splitlines() if l.strip()]
        files = {os.path.basename(x["file"]): x for x in lines if "file" in x}
        summary = next((x for x in lines if x.get("summary")), {})
        return r, files, summary


class ExtractionTest(CensusCase):
    def test_every_part_of_a_docx_counts_including_the_ones_pandoc_misses(self):
        # FA-04's fixture, grown: a marker in the body (split across runs), a table cell, the header, the footer, a
        # text box, a footnote, a comment, a tracked deletion and the document title -- pandoc -t plain found 3 of 5
        body = (para("The member used 1,384 units ", "[DATA ", "NEEDED]", " in the period.")
                + "<w:tbl><w:tr><w:tc>" + para("[REVIEWER TO CONFIRM: method]") + "</w:tc></w:tr></w:tbl>"
                + textbox(para("Summary: [OWNER VERIFY: authorized figure]"))
                + "<w:p><w:del><w:r><w:delText>[TODO: cut]</w:delText></w:r></w:del></w:p>")
        f = docx(self.p("a.docx"), body=body, header=para("Member ID [DATA NEEDED: member id]"),
                 footer=para("Page 1 [FLAG]"), footnotes=para("See the letter [TBD]."),
                 comments=para("[OWNER DECISION NEEDED: send?]"), title="Appeal [DRAFT]")
        r, files, s = self.census(f)
        self.assertEqual(r.returncode, 1, r.stderr)
        row = files["a.docx"]
        self.assertEqual(row["status"], "MARKED")
        parts = row["parts"]
        for part in ("word/document.xml", "word/header1.xml", "word/footer1.xml", "word/footnotes.xml",
                     "word/comments.xml", "docProps/core.xml"):
            self.assertIn(part, parts, part)
        self.assertEqual(parts["word/document.xml"]["gated"], 4)   # body (joined runs), table, text box once, deletion
        self.assertEqual(parts["docProps/core.xml"]["catchall"], 1)          # "[DRAFT]" in the title
        self.assertEqual(row["gated"], 8)
        self.assertEqual((s["exit"], s["mode"], s["marked"]), (1, "strict", 1))

    def test_a_text_box_fallback_copy_counts_only_when_it_differs(self):
        twin = docx(self.p("twin.docx"), body=textbox(para("[DATA NEEDED]")))
        odd = docx(self.p("odd.docx"), body=textbox(para("clean text"), fallback=para("[DATA NEEDED]")))
        _, files, _ = self.census(twin, odd)
        self.assertEqual(files["twin.docx"]["gated"], 1)                     # an exact twin is not counted twice
        self.assertEqual(files["odd.docx"]["gated"], 1)                      # a differing fallback copy is read

    def test_leftover_pointers_are_counted_case_insensitive_and_unclosed(self):
        body = (para('Units denied: 2,400 ', '[src: denial p.2 "Total Units: Denied 2,400"]', '.')
                + para("Utilization was 69.2% [CALC: corp/receipts/util.md] in the period.")
                + para("A figure [ src : letter p.1 \"Approved 240 [sic]\"] and an unclosed [calc: r.md")
                + para("A future grammar [ref: x] stays visible."))
        r, files, s = self.census(docx(self.p("ptr.docx"), body=body))
        row = files["ptr.docx"]
        self.assertEqual((row["pointers"], row["pointer_like"], row["catchall"]), (4, 1, 0), row["hits"])
        self.assertEqual(r.returncode, 1)
        self.assertEqual(s["pointers"], 4)

    def test_a_clean_file_passes_and_text_files_are_read(self):
        clean = docx(self.p("clean.docx"), body=para("Units approved: 240 of 2,640 requested."),
                     header=para("Acme Org"), footer=para("Page 1"))
        md = self.p("notes.md")
        with open(md, "w") as fh:
            fh.write("# Appeal\nThe member used 1,384 units [calc: r.md] and it is TBD.\n")
        r, files, s = self.census(clean, md)
        self.assertEqual(files["clean.docx"]["status"], "CLEAN")
        self.assertEqual((files["notes.md"]["pointers"], files["notes.md"]["gated"]), (1, 1))
        self.assertEqual((r.returncode, s["clean"], s["marked"]), (1, 1, 1))
        r, files, s = self.census(clean)
        self.assertEqual((r.returncode, s["exit"], s["files"]), (0, 0, 1))
        self.assertEqual(len(files["clean.docx"]["md5"]), 32)                  # §GATE ties the receipt to the bytes

    def test_for_signature_keeps_her_flags_but_never_a_pointer(self):
        flags = docx(self.p("draft.docx"), body=para("Method: [REVIEWER TO CONFIRM: measurement method] [CONFIRM]"))
        r, _, s = self.census("--for-signature", flags)
        self.assertEqual((r.returncode, s["mode"], s["gated"], s["catchall"]), (0, "for-signature", 1, 1))
        r, _, s = self.census(flags)
        self.assertEqual(r.returncode, 1)                                      # strict: anything a third party gets
        ptr = docx(self.p("ptr.docx"), body=para('Hours [src: letter p.1 "20 hours"] [REVIEWER TO CONFIRM]'))
        r, _, s = self.census("--for-signature", ptr)
        self.assertEqual((r.returncode, s["pointers"]), (1, 1))

    def test_a_folder_skips_dot_files_and_lockfiles_and_a_scan_or_image_is_unchecked(self):
        d = self.p("SEND TO")
        os.makedirs(d)
        docx(os.path.join(d, "01 Appeal.docx"), body=para("Clean."))
        with open(os.path.join(d, ".corp-deployed.json"), "w") as fh:
            fh.write('{"note": "[src: never counted here]"}')
        with open(os.path.join(d, "~$Appeal.docx"), "w") as fh:
            fh.write("lock")
        with open(os.path.join(d, "scan.png"), "wb") as fh:
            fh.write(b"\x89PNG\r\n\x1a\n")
        r, files, s = self.census(d)
        self.assertEqual(sorted(files), ["01 Appeal.docx", "scan.png"])
        self.assertEqual((files["scan.png"]["status"], r.returncode, s["unchecked"]), ("UNCHECKED", 1, 1))
        r, _, s = self.census("--allow-unchecked", d)
        self.assertEqual((r.returncode, s["allow_unchecked"]), (0, True))

    @unittest.skipUnless(shutil.which("pdftotext"), "needs pdftotext (poppler)")
    def test_a_pdf_is_read_page_by_page_and_a_scan_is_unchecked(self):
        pdf = make_pdf(self.p("letter.pdf"), ["Page one text.", "Page two [DATA NEEDED] here."])
        r, files, s = self.census(pdf)
        row = files["letter.pdf"]
        self.assertEqual((row["kind"], row["gated"], r.returncode), ("pdf", 1, 1))
        self.assertEqual(list(row["parts"]), ["page 2"])
        scan = make_pdf(self.p("scan.pdf"), ["", ""])                          # no text layer at all
        r, files, s = self.census(scan)
        self.assertEqual((files["scan.pdf"]["status"], r.returncode), ("UNCHECKED", 1))
        self.assertIn("no text layer", files["scan.pdf"]["unchecked"])

    def test_bad_input_exits_2(self):
        for args in ([self.p("missing.docx")], ["--strip", self.p("x.md")], ["--stamp", "x", self.p("x.md")]):
            with open(self.p("x.md"), "w") as fh:
                fh.write("x")
            r, _, _ = self.census(*args)
            self.assertEqual(r.returncode, 2, (args, r.stderr))
        bad = self.p("bad.docx")
        with open(bad, "wb") as fh:
            fh.write(b"not a zip")
        self.assertEqual(self.census(bad)[0].returncode, 2)


class EveryPartTest(CensusCase):
    """CEO decision F4 (re-gate FABLE-A RG-02, RG-05; OPUS RG-03; FABLE-B RG-02): every OOXML part, any namespace."""

    def body(self, text):
        return f'<?xml version="1.0"?><w:document {W}><w:body>{para(text)}</w:body></w:document>'

    def test_chart_and_smartart_text_counts_and_a_cached_label_is_never_glued(self):
        # FABLE-A's chart.docx and OPUS's chart.docx / smartart.docx: DrawingML <a:t> and a chart's <c:v> were CLEAN
        chart = (f'<?xml version="1.0"?><c:chartSpace {C} {A}><c:chart><c:title><c:tx><c:rich><a:p><a:r>'
                 '<a:t>10003 utilization [DATA </a:t></a:r><a:r><a:t>NEEDED: authorized figure]</a:t></a:r></a:p>'
                 '</c:rich></c:tx></c:title><c:plotArea><c:barChart><c:ser><c:cat><c:strRef><c:f>Sheet1!$A$2</c:f>'
                 '<c:strCache><c:pt idx="0"><c:v>TBD</c:v></c:pt></c:strCache></c:strRef></c:cat></c:ser>'
                 '</c:barChart></c:plotArea></c:chart></c:chartSpace>')
        smart = ('<?xml version="1.0"?><dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/'
                 f'diagram" {A}><dgm:ptLst><dgm:pt><dgm:t><a:p><a:r><a:t>[REVIEWER TO CONFIRM: method]</a:t></a:r></a:p>'
                 '</dgm:t></dgm:pt></dgm:ptLst></dgm:dataModel>')
        f = package(self.p("chart.docx"), {"word/document.xml": self.body("Utilization graph:"),
                                            "word/charts/chart1.xml": chart, "word/diagrams/data1.xml": smart})
        r, files, s = self.census(f)
        row = files["chart.docx"]
        self.assertEqual((r.returncode, row["status"], row["gated"]), (1, "MARKED", 3), row)
        self.assertEqual(row["parts"]["word/charts/chart1.xml"]["gated"], 2)      # the title (runs joined) + "TBD"
        self.assertEqual(row["parts"]["word/diagrams/data1.xml"]["gated"], 1)
        self.assertIn("[DATA NEEDED: authorized figure]", [h["text"] for h in row["hits"]])

    def test_an_embedded_workbook_is_read_and_an_ole_object_is_unchecked(self):
        # OPUS's embedded.docx: a workbook cell "10003 authorized [REVIEWER TO CONFIRM]" was CLEAN
        xlsx = package_bytes({"xl/sharedStrings.xml": '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/'
                                                      '2006/main"><si><t>10003 authorized [REVIEWER TO CONFIRM]</t></si></sst>'})
        f = package(self.p("embedded.docx"), {"word/document.xml": self.body("See the utilization table below."),
                                               "word/embeddings/Microsoft_Excel_Worksheet.xlsx": xlsx})
        r, files, _ = self.census(f)
        row = files["embedded.docx"]
        self.assertEqual((r.returncode, row["status"], row["gated"]), (1, "MARKED", 1), row)
        self.assertEqual(list(row["parts"]), ["word/embeddings/Microsoft_Excel_Worksheet.xlsx!xl/sharedStrings.xml"])
        ole = package(self.p("ole.docx"), {"word/document.xml": self.body("Clean."),
                                            "word/embeddings/oleObject1.bin": b"\xd0\xcf\x11\xe0 OLE"})
        r, files, s = self.census(ole)
        self.assertEqual((r.returncode, files["ole.docx"]["status"], s["unchecked"]), (1, "UNCHECKED", 1))
        self.assertIn("oleObject1.bin", files["ole.docx"]["unchecked"])
        r, _, s = self.census("--allow-unchecked", ole)
        self.assertEqual((r.returncode, s["allow_unchecked"]), (0, True))            # read by eye, recorded

    def test_strict_ooxml_reads_like_transitional(self):
        # FABLE-A's and OPUS's strict.docx: CLEAN with text_chars 0 before; the transitional twin was the control
        text = 'Member [DATA NEEDED: member id] units [src: letter p.1 "x"]'
        strict = package(self.p("strict.docx"), {"word/document.xml": f'<?xml version="1.0"?><w:document {STRICT_W}>'
                                                                      f'<w:body>{para(text)}</w:body></w:document>'})
        trans = package(self.p("transitional.docx"), {"word/document.xml": self.body(text)})
        r, files, _ = self.census(strict, trans)
        for name in ("strict.docx", "transitional.docx"):
            row = files[name]
            self.assertEqual((row["status"], row["gated"], row["pointers"]), ("MARKED", 1, 1), (name, row))
            self.assertGreater(row["text_chars"], 0)
        self.assertEqual(r.returncode, 1)

    def test_a_docx_whose_main_part_yields_no_text_is_unchecked(self):
        # FABLE-A RG-05: a package whose text census cannot read was CLEAN; an image-only page is the same case
        img = ('<?xml version="1.0"?><w:document ' + W + '><w:body><w:p><w:r><w:drawing/></w:r></w:p></w:body>'
               '</w:document>')
        f = package(self.p("image-only.docx"), {"word/document.xml": img,
                                                 "docProps/core.xml": '<cp:coreProperties xmlns:cp="x" xmlns:dc="d">'
                                                                      '<dc:title>Letter</dc:title></cp:coreProperties>'})
        r, files, s = self.census(f)
        row = files["image-only.docx"]
        self.assertEqual((r.returncode, row["status"], s["unchecked"]), (1, "UNCHECKED", 1), row)
        self.assertIn("yields no text", row["unchecked"])

    def test_a_non_xml_text_part_is_read_and_a_part_census_cannot_read_is_unchecked(self):
        # closure check P2-b (census/extra_fixtures.txt altchunk.docx): a Word altChunk part word/afchunk.htm holding
        # "[DATA NEEDED: authorized]" was CLEAN, gated 0, rc 0
        html = ('<html><body><p>Units requested: [DATA NEEDED: authorized]</p><img alt="[REVIEWER TO CONFIRM: graph]"/>'
                '<script>var x = "[TODO]";</script></body></html>')
        f = package(self.p("altchunk.docx"), {"word/document.xml": self.body("See the attached section."),
                                               "word/afchunk.htm": html})
        r, files, s = self.census(f)
        row = files["altchunk.docx"]
        self.assertEqual((r.returncode, row["status"]), (1, "MARKED"), row)
        self.assertGreaterEqual(row["gated"], 2, row)
        self.assertIn("word/afchunk.htm", row["parts"])
        self.assertIn("[DATA NEEDED: authorized]", [h["text"] for h in row["hits"]])
        self.assertIn("[REVIEWER TO CONFIRM: graph]", [h["text"] for h in row["hits"]])     # the alt text a reader sees
        # a clean html altChunk stays CLEAN, and an html-entity-escaped marker is still found
        clean = package(self.p("clean-af.docx"), {"word/document.xml": self.body("Body text."),
                                                   "word/afchunk.html": "<html><body><p>Nothing to fill.</p></body></html>"})
        r, files, _ = self.census(clean)
        self.assertEqual((r.returncode, files["clean-af.docx"]["status"]), (0, "CLEAN"), files["clean-af.docx"])
        esc = package(self.p("esc-af.docx"), {"word/document.xml": self.body("Body text."),
                                               "word/afchunk.htm": "<p>units &#91;DATA NEEDED: units&#93;</p>"})
        r, files, _ = self.census(esc)
        self.assertEqual((r.returncode, files["esc-af.docx"]["gated"]), (1, 1))
        # .txt and .rtf parts
        txt = package(self.p("txt-af.docx"), {"word/document.xml": self.body("Body."),
                                               "word/afchunk.txt": "line one\nUnits: [DATA NEEDED: units]\n"})
        r, files, _ = self.census(txt)
        self.assertEqual((r.returncode, files["txt-af.docx"]["gated"]), (1, 1), files["txt-af.docx"])
        rtf = package(self.p("rtf-af.docx"), {"word/document.xml": self.body("Body."),
                                               "word/afchunk.rtf": r"{\rtf1\ansi{\fonttbl{\f0 Arial;}}\f0\fs22 Units: [DATA NEEDED: units]\par}"})
        r, files, _ = self.census(rtf)
        self.assertEqual((r.returncode, files["rtf-af.docx"]["gated"]), (1, 1), files["rtf-af.docx"])
        # a part census cannot read is UNCHECKED, never silently CLEAN: an unknown binary, a non-text .txt, an .mht
        for name, data in (("word/afchunk.dat", b"\x00\x01\x02 opaque"), ("word/afchunk.txt", b"\xff\xfe\x00\x00\x9c"),
                           ("word/afchunk.mht", b"MIME-Version: 1.0\r\n\r\nUnits: [DATA NEEDED: x]")):
            with self.subTest(part=name):
                odd = package(self.p("odd.docx"), {"word/document.xml": self.body("Clean body."), name: data})
                r, files, s = self.census(odd)
                row = files["odd.docx"]
                self.assertEqual((r.returncode, row["status"], s["unchecked"]), (1, "UNCHECKED", 1), row)
                self.assertIn(name, row["unchecked"])
                self.assertIn("not XML", row["unchecked"])
        # what Word itself writes is not a false alarm: images, printer settings, fonts, thumbnails
        word = package(self.p("word-real.docx"), {"word/document.xml": self.body("Clean body."),
                                                   "word/media/image1.png": b"\x89PNG", "word/media/image2.jpeg": b"\xff\xd8",
                                                   "word/printerSettings/printerSettings1.bin": b"\x00\x01",
                                                   "word/fonts/font1.odttf": b"\x00\x02", "docProps/thumbnail.jpeg": b"\xff\xd8",
                                                   "word/media/image3.emf": b"\x01\x00"})
        r, files, _ = self.census(word)
        self.assertEqual((r.returncode, files["word-real.docx"]["status"]), (0, "CLEAN"), files["word-real.docx"])
        # an altChunk that is itself a package is opened like an embedding
        inner = package_bytes({"word/document.xml": self.body("Inner [DATA NEEDED: inner]")})
        nested = package(self.p("nested-af.docx"), {"word/document.xml": self.body("Body."), "word/afchunk.docx": inner})
        r, files, _ = self.census(nested)
        self.assertEqual((r.returncode, files["nested-af.docx"]["gated"]), (1, 1), files["nested-af.docx"])

    def test_text_a_reader_sees_in_attributes_counts(self):
        # alt text, a field instruction and a VML WordArt string are in the file the recipient gets
        body = (f'<?xml version="1.0"?><w:document {NS} xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/'
                'wordprocessingDrawing"><w:body>' + para("Figure 1.")
                + '<w:p><w:r><w:drawing><wp:inline><wp:docPr id="1" name="Picture 1" descr="Graph [DATA NEEDED]"/>'
                '</wp:inline></w:drawing></w:r></w:p>'
                '<w:p><w:fldSimple w:instr=" DOCPROPERTY Owner \\* [TBD] "><w:r><w:t>x</w:t></w:r></w:fldSimple></w:p>'
                '<w:p><w:r><w:pict><v:shape><v:textpath string="[FLAG] draft"/></v:shape></w:pict></w:r></w:p>'
                '</w:body></w:document>')
        r, files, _ = self.census(package(self.p("attrs.docx"), {"word/document.xml": body}))
        row = files["attrs.docx"]
        self.assertEqual((r.returncode, row["gated"]), (1, 3), row["hits"])

    def test_workbook_and_slide_text_join_their_runs_and_keep_units_apart(self):
        sst = ('<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><r><t>[DATA </t></r>'
               '<r><t>NEEDED]</t></r></si><si><t>Units</t></si></sst>')
        wb = ('<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheets>'
              '<sheet name="Units [TBD]" sheetId="1"/></sheets></workbook>')
        comments = ('<comments xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><commentList>'
                    '<comment ref="A1"><text><r><t>check TODO</t></r></text></comment></commentList></comments>')
        x = package(self.p("util.xlsx"), {"xl/sharedStrings.xml": sst, "xl/workbook.xml": wb,
                                           "xl/comments1.xml": comments})
        slide = (f'<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" {A}><p:cSld><p:spTree>'
                 '<p:sp><p:txBody><a:p><a:r><a:t>Baseline [REVIEWER TO </a:t></a:r><a:r><a:t>CONFIRM]</a:t></a:r></a:p>'
                 '</p:txBody></p:sp></p:spTree></p:cSld></p:sld>')
        pp = package(self.p("deck.pptx"), {"ppt/slides/slide1.xml": slide})
        r, files, _ = self.census(x, pp)
        self.assertEqual(files["util.xlsx"]["gated"], 3, files["util.xlsx"]["hits"])   # split string, sheet name, comment
        self.assertEqual(files["deck.pptx"]["gated"], 1, files["deck.pptx"]["hits"])
        self.assertEqual(r.returncode, 1)

    def test_every_gated_marker_is_listed_for_the_owner_queue(self):
        # F10 (re-gate OPUS OM-10): every census hit maps to a owner-queue row by file and part -- so none is cut off
        body = "".join(para(f"Item {i} [REVIEWER TO CONFIRM: item {i}]") for i in range(40)) + \
            "".join(para(f"Note {i} [Emphasis added]") for i in range(30))
        r, files, _ = self.census(docx(self.p("many.docx"), body=body))
        row = files["many.docx"]
        gated = [h for h in row["hits"] if h["class"] == "gated"]
        self.assertEqual((row["gated"], len(gated), row["catchall"]), (40, 40, 30))
        self.assertEqual(len(row["hits"]) - len(gated), 25)
        self.assertEqual(row["hits_omitted"], 5)
        self.assertTrue(all(h["part"] == "word/document.xml" for h in gated))

    def test_a_bracketed_citation_is_not_a_pointer(self):
        # FABLE-B RG-02 (census-edge/refs.md, doi.md): a references list failed the delivery gate with no marker in it
        refs = self.p("refs.md")
        with open(refs, "w") as fh:
            fh.write("References\n1. State program handbook [https://agency.example/handbook/"
                     "ref.pdf]\n2. Contact [mailto:records@example.org] for records.\nSmith et al. (2020) "
                     "[doi:10.1000/xyz.123]; [http://example.org/a]; [ftp://x.org/f]; [isbn: 978-0000000002]; "
                     "[pmid: 12345]; [urn:x:y]\n")
        r, files, s = self.census(refs)
        self.assertEqual((r.returncode, files["refs.md"]["status"], s["pointer_like"]), (0, "CLEAN", 0), files["refs.md"])
        other = self.p("other.md")
        with open(other, "w") as fh:
            fh.write("A future grammar [ref: x] and a bare scheme word [http: x] stay pointer-like.\n")
        r, files, s = self.census(other)
        self.assertEqual((r.returncode, s["pointer_like"]), (1, 2))


class StripTest(CensusCase):
    def test_strip_text_into_a_copy_then_census_it(self):
        md = self.p("appeal.md")
        text = ('Units denied: 2,400 [src: denial p.2 "Total Units: Denied 2,400 [sic]"].\n'
                "Utilization was 69.2% [calc: corp/receipts/util.md] in the period.\n"
                "[src: denial p.1 \"Review Outcome: Modified\"] The outcome was modified.\n"
                "Again 2,400 [src: denial p.2 \"Total Units: Denied 2,400 [sic]\"].\n")
        with open(md, "w") as fh:
            fh.write(text)
        out = self.p("v2")
        r, files, s = self.census("--strip", "--out-dir", out, md)
        self.assertEqual(r.returncode, 0, r.stderr)
        with open(os.path.join(out, "appeal.md")) as fh:
            got = fh.read()
        self.assertEqual(got, "Units denied: 2,400.\nUtilization was 69.2% in the period.\nThe outcome was modified.\n"
                              "Again 2,400.\n")
        with open(md) as fh:
            self.assertEqual(fh.read(), text)                                  # never in place
        self.assertEqual(s["stripped"][md], 4)
        self.assertEqual((s["pointers"], s["exit"]), (0, 0))

    def test_strip_a_docx_across_runs_and_parts_with_the_build_stamp(self):
        try:
            import docx as python_docx
        except ImportError:
            self.skipTest("python-docx is not importable")
        d = python_docx.Document()
        p = d.add_paragraph("Units approved: 240 ")
        p.add_run('[src: letter p.2 "Total Units: ')
        p.add_run('Approved 240"]')
        p.add_run(" per the letter.")
        d.add_table(rows=1, cols=1).cell(0, 0).text = "Hours [calc: receipts/hours.md] weekly"
        d.sections[0].header.paragraphs[0].text = "Member [src: letter p.1 \"Member ID 123\"]"
        src = self.p("appeal.docx")
        d.save(src)
        plain = self.p("cover.md")
        with open(plain, "w") as fh:
            fh.write("No pointers here.\n")
        out = self.p("v2")
        r, files, s = self.census("--strip", "--out-dir", out, "--stamp", "Acme Org appeal build", src, plain)
        self.assertEqual(r.returncode, 0, r.stderr + r.stdout)
        self.assertEqual(s["pointers"], 0)
        self.assertEqual(s["stripped"][src], 3)
        self.assertEqual(s["stripped"][plain], 0)
        self.assertTrue(os.path.exists(os.path.join(out, "cover.md")))        # the whole set lands in --out-dir
        got = python_docx.Document(os.path.join(out, "appeal.docx"))
        self.assertEqual(got.paragraphs[0].text, "Units approved: 240 per the letter.")
        self.assertEqual(got.tables[0].cell(0, 0).text, "Hours weekly")
        self.assertEqual(got.sections[0].header.paragraphs[0].text, "Member")
        self.assertEqual(got.core_properties.last_modified_by, "Acme Org appeal build")

    def test_strip_refuses_a_pdf_and_in_place(self):
        pdf = self.p("x.pdf")
        with open(pdf, "wb") as fh:
            fh.write(b"%PDF-1.4\n")
        self.assertEqual(self.census("--strip", "--out-dir", self.p("v2"), pdf)[0].returncode, 2)
        md = self.p("x.md")
        with open(md, "w") as fh:
            fh.write("a [calc: r] b\n")
        self.assertEqual(self.census("--strip", "--out-dir", self.tmp, md)[0].returncode, 2)


if __name__ == "__main__":
    unittest.main()
